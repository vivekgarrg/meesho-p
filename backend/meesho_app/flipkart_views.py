"""
Flipkart, as a marketplace of its own next to Meesho.

Meesho splits its data into orders (what was ordered, and how it ended) and
payments (what was actually paid out); Flipkart gets the same two halves here:

  * FlipkartOrder        — the seller-hub *Orders report*: every order item and
                           its status, settled or not.
  * FlipkartOrderPayment — the *Settlement Report*'s "Orders" sheet: what
                           Flipkart paid for each item, one row per NEFT batch.
  * FlipkartOtherCharge  — the Settlement Report's non-order sheets (ads, fines,
                           SPF, rebates...), which also move bank money.

Profit uses the same mechanics as Meesho — the shared SKU pricing table
(FinalPrice / ParentItemPrice), the business's packaging policy and the
loss-recovery buffer — via views._estimated_profit_pricing_lookup. A delivered
item's cost is item price x qty + tax + packaging + buffer; a returned item
carries no cost and only counts when its settlement is a real loss. Ads, fines
and the other settlement charges are then taken off, the way Meesho takes off
ads and recoveries.
"""

import re
from collections import Counter, defaultdict
from decimal import Decimal

import openpyxl
from django.db import transaction
from django.db.models import Count, Q as DQ, Sum
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from .models import (
    BusinessCostSetting, FlipkartOrder, FlipkartOrderPayment, FlipkartOtherCharge,
)
from .permissions import get_authorized_business
from .views import (
    _FLIPKART_PAYMENT_COMPONENTS, _compute_flipkart_profit_summary, _cost_setting,
    _estimated_profit_pricing_lookup, _flipkart_parent_link_map, _has_unit_price,
    _packaging_charge, _sku_key, _to_num, safe_date, safe_decimal, safe_str,
)

ZERO = Decimal("0")

# Order-report statuses, as Flipkart spells them.
STATUS_DELIVERED = "DELIVERED"
STATUS_CANCELLED = "CANCELLED"
STATUS_RETURNED = "RETURNED"
STATUS_RETURN_REQUESTED = "RETURN_REQUESTED"
_RETURN_STATUSES = {STATUS_RETURNED, STATUS_RETURN_REQUESTED}

# Settlement-report return_type values.
RETURN_LOGISTICS = "LOGISTICS RETURN"   # courier brought it back undelivered (RTO)

PAGE_SIZE_MAX = 200


# ── value cleaning ────────────────────────────────────────────────────────────

def clean_item_id(value):
    """'OI:438512424344074100' -> '438512424344074100' (payments carry no prefix)."""
    text = safe_str(value) or ""
    return text[3:] if text.upper().startswith("OI:") else text


def clean_sku(value):
    """The orders export wraps SKUs as \"\"\"SKU:abc\"\"\"; strip down to `abc`."""
    text = (safe_str(value) or "").strip("\"' ")
    if text.upper().startswith("SKU:"):
        text = text[4:]
    return text.strip("\"' ")


def clean_text(value):
    return (safe_str(value) or "").strip("\"' ")


def _date(value):
    """safe_date, but blank cells give None — safe_date('') hands back pandas' NaT, which the DB rejects."""
    parsed = safe_date(value)
    return parsed if parsed is not None and parsed == parsed else None


def _flag(value):
    return (safe_str(value) or "").upper() == "Y"


def _norm_header(value):
    """First line of a header cell, lower-cased — cells wrap a formula note onto line two."""
    return str(value).split("\n")[0].strip().lower() if value is not None else ""


def _int_or(value, default=1):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _date_params(request):
    return request.GET.get("date_from", "") or None, request.GET.get("date_to", "") or None


def _page_params(request, default_size=50):
    try:
        page = max(1, int(request.GET.get("page", 1)))
        size = min(PAGE_SIZE_MAX, max(1, int(request.GET.get("page_size", default_size))))
    except ValueError:
        page, size = 1, default_size
    return page, size


# ── Orders report upload ──────────────────────────────────────────────────────

@api_view(["POST"])
@parser_classes([MultiPartParser])
def flipkart_orders_upload(request, business_id):
    """
    Upload Flipkart's Orders report (.xlsx) and save every order item from its
    "Orders" sheet to FlipkartOrder, keyed by (business, order_item_id) — so
    uploading a later report refreshes an item's status (SHIPPED -> DELIVERED
    -> RETURNED) in place rather than duplicating it.
    """
    business = get_authorized_business(request, business_id)
    file = request.FILES.get("file")
    if not file:
        return Response({"error": "No file provided"}, status=status.HTTP_400_BAD_REQUEST)

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception as e:
        return Response({"error": f"Unable to read Excel file: {e}"}, status=status.HTTP_400_BAD_REQUEST)

    if "Orders" not in wb.sheetnames:
        return Response(
            {"error": "This file has no 'Orders' sheet — is it a Flipkart Orders report?"},
            status=status.HTTP_400_BAD_REQUEST,
        )

    rows = wb["Orders"].iter_rows(values_only=True)
    header = next(rows, None) or ()
    col = {}
    for i, name in enumerate(header):
        col.setdefault(_norm_header(name), i)
    if not {"order_item_id", "order_id", "order_item_status", "sku"}.issubset(col):
        return Response(
            {"error": "Could not find the expected columns (order_item_id, order_id, "
                      "order_item_status, sku) in the 'Orders' sheet — is this the "
                      "Settlement Report? Upload that under Flipkart payments."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    def get(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else None

    total_rows = invalid_rows = created = updated = 0
    with transaction.atomic():
        for row in rows:
            order_item_id = clean_item_id(get(row, "order_item_id"))
            if not order_item_id:
                continue
            total_rows += 1
            order_id = safe_str(get(row, "order_id")) or ""
            if not order_id:
                invalid_rows += 1
                continue

            _, was_created = FlipkartOrder.objects.update_or_create(
                business=business,
                order_item_id=order_item_id,
                defaults={
                    "order_id": order_id,
                    "order_date": _date(get(row, "order_date")),
                    "order_approval_date": _date(get(row, "order_approval_date")),
                    "order_item_status": (safe_str(get(row, "order_item_status")) or "").upper(),
                    "sku": clean_sku(get(row, "sku")),
                    "fsn": safe_str(get(row, "fsn")) or "",
                    "product_title": clean_text(get(row, "product_title"))[:500],
                    "quantity": max(1, _int_or(get(row, "quantity"))),
                    "fulfilment_type": safe_str(get(row, "fulfilment_type")) or "",
                    "logistics_partner": (safe_str(get(row, "delivery_logistics_partner"))
                                          or safe_str(get(row, "pickup_logistics_partner")) or ""),
                    "delivery_tracking_id": safe_str(get(row, "delivery_tracking_id")) or "",
                    "cancellation_date": _date(get(row, "order_cancellation_date")),
                    "cancellation_reason": safe_str(get(row, "cancellation_reason")) or "",
                    "return_approval_date": _date(get(row, "order_return_approval_date")),
                    "return_id": safe_str(get(row, "return_id")) or "",
                    "return_reason": safe_str(get(row, "return_reason")) or "",
                    "return_sub_reason": safe_str(get(row, "return_sub_reason")) or "",
                    "dispatch_by_date": _date(get(row, "dispatch_by_date")),
                    "dispatched_date": _date(get(row, "dispatched_date")),
                    "dispatch_sla_breached": _flag(get(row, "dispatch_sla_breached")),
                    "deliver_by_date": _date(get(row, "deliver_by_date")),
                    "delivery_date": _date(get(row, "order_delivery_date")),
                    "delivery_sla_breached": _flag(get(row, "delivery_sla_breached")),
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1

    return Response({
        "total_rows": total_rows,
        "invalid_rows": invalid_rows,
        "created": created,
        "updated": updated,
    })


# ── Settlement Report: the non-order sheets ───────────────────────────────────

# sheet -> (category, header of the id column, header(s) describing the row).
# Every one of these sheets carries "NEFT ID", a payment date and a
# "Settlement Value" column; only the reference and wording differ.
_OTHER_SHEETS = {
    "Ads":                          (FlipkartOtherCharge.CATEGORY_ADS, "campaign / transaction id", ("type",)),
    "Fines":                        (FlipkartOtherCharge.CATEGORY_FINES, "transaction id", ("transaction details",)),
    "Non_Order_SPF":                (FlipkartOtherCharge.CATEGORY_SPF, "claim id", ("protection reason",)),
    "MP Fee Rebate":                (FlipkartOtherCharge.CATEGORY_MP_FEE_REBATE, "order item id", ("sku",)),
    "Storage_Recall":               (FlipkartOtherCharge.CATEGORY_STORAGE_RECALL, "recall id", ("service name",)),
    "Value Added Services":         (FlipkartOtherCharge.CATEGORY_SERVICES, "service order id", ("service name",)),
    "Support Services Services":    (FlipkartOtherCharge.CATEGORY_SERVICES, "service order id", ("service name",)),
    "Google Ads Services":          (FlipkartOtherCharge.CATEGORY_SERVICES, "service order id", ("service name",)),
    "Review Accelerator Services":  (FlipkartOtherCharge.CATEGORY_SERVICES, "service order id", ("service name",)),
    "Insight Subscription Services": (FlipkartOtherCharge.CATEGORY_SERVICES, "service order id", ("service name",)),
    "TCS_Recovery":                 (FlipkartOtherCharge.CATEGORY_TAX, "service order id", ("settlement type",)),
    "TDS":                          (FlipkartOtherCharge.CATEGORY_TAX, "id", ()),
}


def import_flipkart_other_charges(wb, business):
    """
    Save the non-order sheets of a Settlement Report workbook as
    FlipkartOtherCharge rows. Re-importing the same report updates in place.
    Returns {"created": n, "updated": n} counting only rows with a non-zero
    settlement value (an ads "redeem" row moves wallet balance, not cash).
    """
    created = updated = 0
    for sheet, (category, ref_header, desc_headers) in _OTHER_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows[:5]) if r and r[0] == "NEFT ID"), None)
        if header_idx is None:
            continue
        col = {}
        for i, name in enumerate(rows[header_idx]):
            col.setdefault(_norm_header(name), i)
        amount_col = next((i for name, i in col.items() if name.startswith("settlement value")), None)
        if amount_col is None:
            continue

        def get(row, key):
            i = col.get(key)
            return row[i] if i is not None and i < len(row) else None

        for row in rows[header_idx + 1:]:
            amount = safe_decimal(row[amount_col] if amount_col < len(row) else None) or ZERO
            neft_id = safe_str(row[0]) if row else None
            if not neft_id or amount == ZERO:
                continue
            description = " ".join(filter(None, (safe_str(get(row, h)) for h in desc_headers)))[:255]
            _, was_created = FlipkartOtherCharge.objects.update_or_create(
                business=business,
                source_sheet=sheet,
                neft_id=neft_id,
                reference_id=(safe_str(get(row, ref_header)) or "")[:200],
                description=description,
                defaults={
                    "category": category,
                    "payment_date": _date(get(row, "payment date")),
                    "amount": amount,
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1
    return {"created": created, "updated": updated}


def _other_charges_summary(business, date_from, date_to):
    """Settlement-report charges in the period (by payment date), grouped by category."""
    qs = FlipkartOtherCharge.objects.filter(business=business)
    if date_from:
        qs = qs.filter(payment_date__gte=date_from)
    if date_to:
        qs = qs.filter(payment_date__lte=date_to)
    labels = dict(FlipkartOtherCharge.CATEGORY_CHOICES)
    by_category = [
        {"category": r["category"], "label": labels.get(r["category"], r["category"]),
         "count": r["count"], "total": _to_num(r["total"] or ZERO)}
        for r in qs.values("category").annotate(count=Count("id"), total=Sum("amount")).order_by("total")
    ]
    total = qs.aggregate(t=Sum("amount"))["t"] or ZERO
    return {"total": _to_num(total), "count": qs.count(), "by_category": by_category}


@api_view(["GET"])
def flipkart_other_charges_list(request, business_id):
    """The non-order settlement lines (ads top-ups, fines, ...) — newest first."""
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)
    category = request.GET.get("category", "").strip()
    qs = FlipkartOtherCharge.objects.filter(business=business)
    if date_from:
        qs = qs.filter(payment_date__gte=date_from)
    if date_to:
        qs = qs.filter(payment_date__lte=date_to)
    if category:
        qs = qs.filter(category=category)
    page, size = _page_params(request)
    total = qs.count()
    labels = dict(FlipkartOtherCharge.CATEGORY_CHOICES)
    results = [
        {
            "id": r.id, "category": r.category, "category_label": labels.get(r.category, r.category),
            "source_sheet": r.source_sheet, "neft_id": r.neft_id,
            "payment_date": str(r.payment_date) if r.payment_date else None,
            "reference_id": r.reference_id, "description": r.description, "amount": _to_num(r.amount),
        }
        for r in qs[(page - 1) * size: page * size]
    ]
    return Response({"total": total, "page": page, "page_size": size,
                     "summary": _other_charges_summary(business, date_from, date_to), "results": results})


# ── The P&L of one order item ─────────────────────────────────────────────────

def _is_return(return_type):
    return bool(return_type) and return_type.strip().upper() != "NA"


class FlipkartPricing:
    """The SKU pricing + packaging policy a Flipkart P&L is costed with — built once per request."""

    def __init__(self, business):
        self.get_price, self.known_skus, self.canonical_sku = _estimated_profit_pricing_lookup(business)
        self.cost_setting = _cost_setting(business)

    def item_pnl(self, rows):
        """
        P&L of one order item from all its settlement rows (the sale and, later,
        its return reversal are separate rows of the same order_item_id — they
        are instalments of one unit and are costed once, never once per row).

        Same rules as _compute_flipkart_profit_summary:
          * no unit price on the order date -> left out ("missing_price")
          * returned, settlement >= 0       -> left out ("non_loss_return")
          * returned, settlement <  0       -> counts at the settlement, no cost
          * otherwise                       -> settlement - (item cost x qty +
                                               tax + packaging + buffer)
        """
        primary = max(rows, key=lambda r: abs(r.sale_amount or ZERO))
        sku = primary.seller_sku or next((r.seller_sku for r in rows if r.seller_sku), "") or ""
        settlement = sum((r.settlement_value or ZERO for r in rows), ZERO)
        returned = any(_is_return(r.return_type) for r in rows)
        order_date = primary.order_date or next((r.order_date for r in rows if r.order_date), None)
        quantity = primary.quantity or next((r.quantity for r in rows if r.quantity), 1)

        cost, net, excluded = ZERO, settlement, None
        if _sku_key(sku) not in self.known_skus or not _has_unit_price(self.get_price(sku, order_date)[0]):
            excluded, net = "missing_price", None
        elif returned:
            if settlement >= 0:
                excluded, net = "non_loss_return", ZERO
        else:
            item_price, packaging_cost, tax_pct = self.get_price(sku, order_date)
            qty = Decimal(str(quantity or 1))
            purchase = Decimal(str(item_price or 0)) * qty
            tax = purchase * Decimal(str(tax_pct or 0)) / Decimal("100")
            packaging = _packaging_charge(BusinessCostSetting.DELIVERED, Decimal(str(packaging_cost or 0)), self.cost_setting)
            buffer_cost = Decimal(str(self.get_price.buffer_for(sku) or 0)) * qty
            cost = purchase + tax + packaging + buffer_cost
            net = settlement - cost

        return {
            "sku": sku, "quantity": quantity, "order_date": order_date, "returned": returned,
            "return_type": next((r.return_type for r in rows if _is_return(r.return_type)), "NA"),
            "settlement": settlement, "cost": cost, "net": net, "excluded_reason": excluded,
            "primary": primary,
        }


def _payment_row(r):
    return {
        "neft_id": r.neft_id, "payment_date": str(r.payment_date) if r.payment_date else None,
        "return_type": r.return_type, "item_return_status": r.item_return_status,
        **{f: _to_num(getattr(r, f) or ZERO) for f in _FLIPKART_PAYMENT_COMPONENTS},
    }


def _group_status(item_pnls):
    """DELIVERED / RETURN (customer) / RTO (courier brought it back) for a whole order."""
    types = {p["return_type"].strip().upper() for p in item_pnls if p["returned"]}
    if RETURN_LOGISTICS in types:
        return "RTO"
    if types:
        return "RETURN"
    return "DELIVERED"


def _filter_payments(business, date_from, date_to):
    qs = FlipkartOrderPayment.objects.filter(business=business)
    if date_from:
        qs = qs.filter(order_date__gte=date_from)
    if date_to:
        qs = qs.filter(order_date__lte=date_to)
    return qs


@api_view(["GET"])
def flipkart_payments_grouped(request, business_id):
    """
    Settlement payments grouped by Flipkart order id — the Flipkart twin of
    Meesho's /orders/grouped/ — with each order's profit worked out the same
    way the P&L is (see FlipkartPricing.item_pnl).

    Params: date_from / date_to (by order date), status = DELIVERED | RETURN |
    RTO, search (order id, order item id, SKU, product title), page, page_size.
    """
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)
    status_f = request.GET.get("status", "").strip().upper()
    search = request.GET.get("search", "").strip()
    page, size = _page_params(request)

    qs = _filter_payments(business, date_from, date_to)
    if search:
        title_items = FlipkartOrder.objects.filter(
            business=business, product_title__icontains=search).values("order_item_id")
        qs = qs.filter(
            DQ(order_id__icontains=search) | DQ(order_item_id__icontains=search)
            | DQ(seller_sku__icontains=search) | DQ(order_item_id__in=title_items))

    rows = list(qs.order_by("order_id", "payment_date", "id"))
    report = {o.order_item_id: o for o in FlipkartOrder.objects.filter(
        business=business, order_item_id__in={r.order_item_id for r in rows})}
    pricing = FlipkartPricing(business)

    by_order = defaultdict(lambda: defaultdict(list))
    for r in rows:
        by_order[r.order_id][r.order_item_id].append(r)

    groups = []
    for order_id, items in by_order.items():
        pnls = [pricing.item_pnl(item_rows) for item_rows in items.values()]
        counted = [p for p in pnls if p["net"] is not None]
        group_dates = [p["order_date"] for p in pnls if p["order_date"]]
        groups.append({
            "order_id": order_id,
            "status": _group_status(pnls),
            "order_date": str(min(group_dates)) if group_dates else None,
            "item_count": len(pnls),
            "row_count": sum(len(v) for v in items.values()),
            "settlement": _to_num(sum((p["settlement"] for p in pnls), ZERO)),
            "cost": _to_num(sum((p["cost"] for p in pnls), ZERO)),
            # None when none of the order's items could be costed (no unit price).
            "net_profit": _to_num(sum((p["net"] for p in counted), ZERO)) if counted else None,
            "missing_price": any(p["excluded_reason"] == "missing_price" for p in pnls),
            "_items": list(zip(items.items(), pnls)),
        })

    status_counts = Counter(g["status"] for g in groups)
    status_totals = defaultdict(lambda: {"settlement": 0.0, "profit": 0.0})
    for g in groups:
        status_totals[g["status"]]["settlement"] += g["settlement"]
        status_totals[g["status"]]["profit"] += g["net_profit"] or 0.0
    kpi = {
        "total_groups": len(groups),
        "total_settlement": round(sum(g["settlement"] for g in groups), 2),
        "total_cost": round(sum(g["cost"] for g in groups), 2),
        "total_profit": round(sum(g["net_profit"] or 0.0 for g in groups), 2),
        "missing_price_orders": sum(1 for g in groups if g["missing_price"]),
        "by_status": [
            {"status": s, "count": status_counts[s], "settlement": round(status_totals[s]["settlement"], 2),
             "profit": round(status_totals[s]["profit"], 2)}
            for s in ("DELIVERED", "RETURN", "RTO") if status_counts.get(s)
        ],
    }

    if status_f:
        groups = [g for g in groups if g["status"] == status_f]
    groups.sort(key=lambda g: g["order_date"] or "", reverse=True)
    total = len(groups)
    page_groups = groups[(page - 1) * size: page * size]

    for g in page_groups:
        items_out = []
        for (item_id, item_rows), p in g["_items"]:
            ordered = report.get(item_id)
            item_rows = sorted(item_rows, key=lambda r: (r.payment_date is None, r.payment_date))
            items_out.append({
                "order_item_id": item_id,
                "sku": p["sku"],
                "product_title": ordered.product_title if ordered else "",
                "order_report_status": ordered.order_item_status if ordered else None,
                "quantity": p["quantity"],
                "is_return": p["returned"],
                "return_type": p["return_type"],
                "settlement": _to_num(p["settlement"]),
                "cost": _to_num(p["cost"]),
                "net_profit": _to_num(p["net"]) if p["net"] is not None else None,
                "excluded_reason": p["excluded_reason"],
                "rows": [_payment_row(r) for r in item_rows],
            })
        g["items"] = items_out
        first = items_out[0] if items_out else {}
        g["sku"] = first.get("sku", "")
        g["product_title"] = first.get("product_title", "")
        del g["_items"]

    return Response({"total": total, "page": page, "page_size": size, "kpi": kpi, "groups": page_groups})


# ── Orders report: list + analytics ──────────────────────────────────────────

def _settled_item_ids(business):
    return set(FlipkartOrderPayment.objects.filter(business=business).values_list("order_item_id", flat=True))


def _payment_status(return_type):
    """A Flipkart settlement row's outcome, in the same words the Meesho
    payments chart uses."""
    rt = (return_type or "").strip().lower()
    if not rt or rt == "na":
        return "Delivered"
    if "courier" in rt or "rto" in rt:
        return "RTO"
    if "customer" in rt or "return" in rt:
        return "Return"
    return rt.replace("_", " ").title()


def _payment_analytics(business, date_from, date_to):
    """The payments half of the Overview charts — the same shape the Meesho
    dashboard sends: settlements per payment day, and payments by settlement
    status. Covers the orders placed in the period, like the profit figures."""
    qs = FlipkartOrderPayment.objects.filter(business=business)
    if date_from:
        qs = qs.filter(order_date__gte=date_from)
    if date_to:
        qs = qs.filter(order_date__lte=date_to)
    daily = (qs.exclude(payment_date=None).values("payment_date")
             .annotate(count=Count("order_item_id", distinct=True), total=Sum("settlement_value"))
             .order_by("payment_date"))
    per_item = {}
    for item_id, return_type, value in qs.values_list("order_item_id", "return_type", "settlement_value"):
        e = per_item.setdefault(item_id, {"status": "Delivered", "total": Decimal("0")})
        st = _payment_status(return_type)
        if st != "Delivered":
            e["status"] = st      # any return row decides the item's outcome
        e["total"] += value or Decimal("0")
    by = defaultdict(lambda: {"count": 0, "total_settlement": Decimal("0")})
    for e in per_item.values():
        by[e["status"]]["count"] += 1
        by[e["status"]]["total_settlement"] += e["total"]
    return {
        "daily": [{"payment_date": str(r["payment_date"]), "count": r["count"], "total": float(r["total"] or 0)} for r in daily],
        "by_status": sorted(({"status": k, "count": v["count"], "total_settlement": float(v["total_settlement"])}
                             for k, v in by.items()), key=lambda r: -r["count"]),
        "settled_items": len(per_item),
    }


def _order_analytics(business, date_from, date_to):
    """What the Orders report says: outcomes, SLA misses and how much has actually been paid."""
    qs = FlipkartOrder.objects.filter(business=business)
    if date_from:
        qs = qs.filter(order_date__gte=date_from)
    if date_to:
        qs = qs.filter(order_date__lte=date_to)

    items = list(qs.only("order_item_id", "order_item_status", "quantity", "sku", "order_date",
                         "return_reason", "cancellation_reason", "dispatch_sla_breached",
                         "delivery_sla_breached"))
    payments = defaultdict(list)
    for item_id, return_type in FlipkartOrderPayment.objects.filter(
            business=business, order_item_id__in=[i.order_item_id for i in items]
    ).values_list("order_item_id", "return_type"):
        payments[item_id].append(return_type)

    by_status = Counter(i.order_item_status or "UNKNOWN" for i in items)
    total = len(items)
    cancelled = by_status.get(STATUS_CANCELLED, 0)
    returned = sum(by_status.get(s, 0) for s in _RETURN_STATUSES)
    shipped_out = total - cancelled          # items that left the warehouse (or could have)

    delivered_unsettled = sum(
        1 for i in items if i.order_item_status == STATUS_DELIVERED and i.order_item_id not in payments)
    # Paid as a plain sale, but the report already shows it coming back: the
    # reversal is on its way and that payout will be clawed back.
    returned_after_settlement = sum(
        1 for i in items if i.order_item_status in _RETURN_STATUSES
        and i.order_item_id in payments and not any(_is_return(rt) for rt in payments[i.order_item_id]))

    daily = Counter(str(i.order_date) for i in items if i.order_date)
    return {
        "total_items": total,
        "total_units": sum(i.quantity for i in items),
        "by_status": [{"status": s, "count": n} for s, n in by_status.most_common()],
        "delivered": by_status.get(STATUS_DELIVERED, 0),
        "returned": returned,
        "cancelled": cancelled,
        "return_rate_pct": round(returned * 100.0 / shipped_out, 1) if shipped_out else 0.0,
        "cancel_rate_pct": round(cancelled * 100.0 / total, 1) if total else 0.0,
        "settled_items": sum(1 for i in items if i.order_item_id in payments),
        "delivered_unsettled": delivered_unsettled,
        "returned_after_settlement": returned_after_settlement,
        "dispatch_sla_breached": sum(1 for i in items if i.dispatch_sla_breached),
        "delivery_sla_breached": sum(1 for i in items if i.delivery_sla_breached),
        "return_reasons": [{"reason": r, "count": n} for r, n in Counter(
            i.return_reason for i in items if i.return_reason).most_common(6)],
        "cancel_reasons": [{"reason": r, "count": n} for r, n in Counter(
            i.cancellation_reason for i in items if i.cancellation_reason).most_common(6)],
        "daily": [{"order_date": d, "count": n} for d, n in sorted(daily.items())],
    }


@api_view(["GET"])
def flipkart_orders_list(request, business_id):
    """
    The Orders report, newest first, each item flagged with whether (and for how
    much) it has been settled. Params: date_from / date_to, status, settled =
    yes | no, search (order id, item id, SKU, title), page, page_size.
    """
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)
    status_f = request.GET.get("status", "").strip().upper()
    settled_f = request.GET.get("settled", "").strip().lower()
    search = request.GET.get("search", "").strip()
    page, size = _page_params(request)

    qs = FlipkartOrder.objects.filter(business=business)
    if date_from:
        qs = qs.filter(order_date__gte=date_from)
    if date_to:
        qs = qs.filter(order_date__lte=date_to)
    if status_f:
        qs = qs.filter(order_item_status=status_f)
    if search:
        qs = qs.filter(DQ(order_id__icontains=search) | DQ(order_item_id__icontains=search)
                       | DQ(sku__icontains=search) | DQ(product_title__icontains=search))

    settled_ids = _settled_item_ids(business)
    if settled_f == "yes":
        qs = qs.filter(order_item_id__in=settled_ids)
    elif settled_f == "no":
        qs = qs.exclude(order_item_id__in=settled_ids)

    total = qs.count()
    page_rows = list(qs[(page - 1) * size: page * size])
    paid = defaultdict(lambda: ZERO)
    for item_id, value in FlipkartOrderPayment.objects.filter(
            business=business, order_item_id__in=[r.order_item_id for r in page_rows]
    ).values_list("order_item_id", "settlement_value"):
        paid[item_id] += value or ZERO

    results = [
        {
            "id": r.id, "order_item_id": r.order_item_id, "order_id": r.order_id,
            "order_date": str(r.order_date) if r.order_date else None,
            "status": r.order_item_status, "sku": r.sku, "product_title": r.product_title,
            "quantity": r.quantity, "logistics_partner": r.logistics_partner,
            "delivery_tracking_id": r.delivery_tracking_id,
            "dispatched_date": str(r.dispatched_date) if r.dispatched_date else None,
            "delivery_date": str(r.delivery_date) if r.delivery_date else None,
            "cancellation_reason": r.cancellation_reason,
            "return_reason": r.return_reason, "return_sub_reason": r.return_sub_reason,
            "dispatch_sla_breached": r.dispatch_sla_breached,
            "delivery_sla_breached": r.delivery_sla_breached,
            "settled": r.order_item_id in settled_ids,
            "settlement_value": _to_num(paid[r.order_item_id]) if r.order_item_id in settled_ids else None,
        }
        for r in page_rows
    ]
    return Response({
        "total": total, "page": page, "page_size": size, "results": results,
        "analytics": _order_analytics(business, date_from, date_to),
    })


# ── Overview + SKU analysis ───────────────────────────────────────────────────

@api_view(["GET"])
def flipkart_overview(request, business_id):
    """
    Everything the Overview tab shows for Flipkart: the item P&L (from the
    settlement payments and the SKU pricing table), the settlement-report
    charges taken off it, and what the Orders report says about outcomes.
    Each business sees only its own rows.
    """
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)

    summary = _compute_flipkart_profit_summary(business, date_from, date_to)
    other = _other_charges_summary(business, date_from, date_to)
    order_profit = summary["totals"]["net_profit"]
    return Response({
        "has_data": summary["has_any_rows"] or FlipkartOrder.objects.filter(business=business).exists(),
        "summary": summary,
        "other_charges": other,
        "net_profit": round(order_profit + other["total"], 2),
        "orders": _order_analytics(business, date_from, date_to),
        "payments": _payment_analytics(business, date_from, date_to),
    })


@api_view(["GET"])
def flipkart_overall(request, business_id):
    """
    Lifetime Flipkart totals for the Overview's hidden "till now" card — every
    settled item and every settlement-report charge, whatever period the page is
    filtered to. Cached until data changes (api_cache._LONG_LIVED).
    """
    business = get_authorized_business(request, business_id)
    summary = _compute_flipkart_profit_summary(business, None, None)
    other = _other_charges_summary(business, None, None)
    pay = summary["payment_breakdown"]
    return Response({
        "gross_revenue": pay.get("sale_amount", 0),
        "total_settled": pay.get("settlement_value", 0),
        "net_profit":    round(summary["totals"]["net_profit"] + other["total"], 2),
        "order_count":   summary["totals"]["order_count"],
    })


def _month(d):
    return d.strftime("%Y-%m") if d else "Unknown"


def _days_late(actual, due):
    """Whole days past the deadline; None when either date is missing."""
    return (actual - due).days if actual and due else None


@api_view(["GET"])
def flipkart_deep_dive(request, business_id):
    """
    The numbers behind three Overview figures, taken apart: what Flipkart's
    marketplace fee really is (rate, trend, by fulfilment / category / SKU),
    where SLAs were missed (how late, who, what it cost), and every fine and
    other settlement charge line by line.
    """
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)

    # ── Marketplace fees (settlement rows, by order date) ─────────────────────
    pay_qs = FlipkartOrderPayment.objects.filter(business=business)
    if date_from:
        pay_qs = pay_qs.filter(order_date__gte=date_from)
    if date_to:
        pay_qs = pay_qs.filter(order_date__lte=date_to)

    def fee_of(r):
        # Both are negative when charged and positive when refunded on a return.
        return -float((r["marketplace_fee"] or 0) + (r["gst_on_mp_fees"] or 0))

    def bucket(rows, key):
        out = {}
        for r in rows:
            b = out.setdefault(key(r), {"sale": 0.0, "fee": 0.0, "items": 0, "settlement": 0.0})
            b["sale"] += float(r["sale_amount"] or 0)
            b["fee"] += fee_of(r)
            b["settlement"] += float(r["settlement_value"] or 0)
            b["items"] += 1
        return out

    rows = list(pay_qs.values("order_item_id", "seller_sku", "order_date", "fulfilment_type", "product_sub_category",
                              "sale_amount", "marketplace_fee", "gst_on_mp_fees", "settlement_value",
                              "taxes", "tcs", "tds", "protection_fund"))

    def finish(d, label_key="label", top=None):
        out = []
        for k, b in d.items():
            sale, fee = round(b["sale"], 2), round(b["fee"], 2)
            out.append({label_key: k, "sale": sale, "fee": fee, "items": b["items"],
                        "rate_pct": round(fee * 100.0 / sale, 1) if sale > 0 else None,
                        "fee_per_item": round(fee / b["items"], 2) if b["items"] else 0})
        out.sort(key=lambda x: -x["fee"])
        return out[:top] if top else out

    total_sale = sum(float(r["sale_amount"] or 0) for r in rows)
    total_fee = sum(fee_of(r) for r in rows)
    mp_fee = -sum(float(r["marketplace_fee"] or 0) for r in rows)
    gst_fee = -sum(float(r["gst_on_mp_fees"] or 0) for r in rows)
    fees = {
        "items": len(rows),
        "sale": round(total_sale, 2),
        "marketplace_fee": round(mp_fee, 2),
        "gst_on_fee": round(gst_fee, 2),
        "total_fee": round(total_fee, 2),
        "rate_pct": round(total_fee * 100.0 / total_sale, 1) if total_sale > 0 else None,
        "avg_fee_per_item": round(total_fee / len(rows), 2) if rows else 0,
        "by_month": sorted(finish(bucket(rows, lambda r: _month(r["order_date"])), "month"), key=lambda x: x["month"]),
        "by_fulfilment": finish(bucket(rows, lambda r: r["fulfilment_type"] or "Unknown")),
        "by_category": finish(bucket(rows, lambda r: r["product_sub_category"] or "Unknown"), top=12),
        "by_sku": finish(bucket(rows, lambda r: r["seller_sku"] or "(blank SKU)"), "sku", top=15),
        # Single items where the fee ate the biggest share of the sale — the ones to question.
        "highest_rate_items": sorted(
            [{"order_item_id": r["order_item_id"], "sku": r["seller_sku"], "sale": float(r["sale_amount"]),
              "fee": round(fee_of(r), 2), "rate_pct": round(fee_of(r) * 100.0 / float(r["sale_amount"]), 1)}
             for r in rows if (r["sale_amount"] or 0) > 0 and fee_of(r) > 0],
            key=lambda x: -x["rate_pct"])[:10],
    }

    # ── SLA breaches (orders report) ──────────────────────────────────────────
    o_qs = FlipkartOrder.objects.filter(business=business)
    if date_from:
        o_qs = o_qs.filter(order_date__gte=date_from)
    if date_to:
        o_qs = o_qs.filter(order_date__lte=date_to)
    orders = list(o_qs)
    live = [o for o in orders if o.order_item_status != STATUS_CANCELLED]
    d_late = [o for o in orders if o.dispatch_sla_breached]
    v_late = [o for o in orders if o.delivery_sla_breached]
    breached = [o for o in orders if o.dispatch_sla_breached or o.delivery_sla_breached]

    def out_rate(group):
        if not group:
            return None
        gone = sum(1 for o in group if o.order_item_status in _RETURN_STATUSES or o.order_item_status == STATUS_CANCELLED)
        return round(gone * 100.0 / len(group), 1)

    def sla_group(key):
        out = {}
        for o in orders:
            b = out.setdefault(key(o), {"items": 0, "dispatch": 0, "delivery": 0})
            b["items"] += 1
            b["dispatch"] += 1 if o.dispatch_sla_breached else 0
            b["delivery"] += 1 if o.delivery_sla_breached else 0
        rows_ = [{"label": k, **b, "breached": b["dispatch"] + b["delivery"],
                  "rate_pct": round((b["dispatch"] + b["delivery"]) * 100.0 / b["items"], 1)}
                 for k, b in out.items()]
        return rows_

    breached_ids = {o.pk for o in breached}
    dispatch_days = [x for x in (_days_late(o.dispatched_date, o.dispatch_by_date) for o in d_late) if x is not None and x > 0]
    delivery_days = [x for x in (_days_late(o.delivery_date, o.deliver_by_date) for o in v_late) if x is not None and x > 0]
    by_month = sla_group(lambda o: _month(o.order_date))
    sla = {
        "items": len(orders),
        "dispatch_breached": len(d_late),
        "delivery_breached": len(v_late),
        "any_breached": len(breached),
        "dispatch_rate_pct": round(len(d_late) * 100.0 / len(live), 1) if live else 0.0,
        "delivery_rate_pct": round(len(v_late) * 100.0 / len(live), 1) if live else 0.0,
        "avg_dispatch_days_late": round(sum(dispatch_days) / len(dispatch_days), 1) if dispatch_days else None,
        "avg_delivery_days_late": round(sum(delivery_days) / len(delivery_days), 1) if delivery_days else None,
        "worst_dispatch_days_late": max(dispatch_days) if dispatch_days else None,
        # Does missing the SLA coincide with the order going wrong?
        "bad_outcome_rate_breached_pct": out_rate(breached),
        "bad_outcome_rate_on_time_pct": out_rate([o for o in orders if o.pk not in breached_ids]),
        "by_month": sorted(by_month, key=lambda x: x["label"]),
        "by_partner": sorted(sla_group(lambda o: o.logistics_partner or "Unknown"), key=lambda x: -x["breached"])[:8],
        "by_fulfilment": sorted(sla_group(lambda o: o.fulfilment_type or "Unknown"), key=lambda x: -x["breached"]),
        "by_sku": sorted(sla_group(lambda o: o.sku or "(blank SKU)"), key=lambda x: -x["breached"])[:10],
        "worst_orders": [
            {"order_item_id": o.order_item_id, "sku": o.sku, "status": o.order_item_status,
             "order_date": str(o.order_date) if o.order_date else None,
             "dispatch_by": str(o.dispatch_by_date) if o.dispatch_by_date else None,
             "dispatched": str(o.dispatched_date) if o.dispatched_date else None,
             "dispatch_days_late": _days_late(o.dispatched_date, o.dispatch_by_date) if o.dispatch_sla_breached else None,
             "deliver_by": str(o.deliver_by_date) if o.deliver_by_date else None,
             "delivered": str(o.delivery_date) if o.delivery_date else None,
             "delivery_days_late": _days_late(o.delivery_date, o.deliver_by_date) if o.delivery_sla_breached else None}
            for o in sorted(breached, key=lambda o: -(max(
                _days_late(o.dispatched_date, o.dispatch_by_date) or 0 if o.dispatch_sla_breached else 0,
                _days_late(o.delivery_date, o.deliver_by_date) or 0 if o.delivery_sla_breached else 0)))[:15]
        ],
    }

    # ── Fines & other settlement charges (by payment date) ────────────────────
    c_qs = FlipkartOtherCharge.objects.filter(business=business)
    if date_from:
        c_qs = c_qs.filter(payment_date__gte=date_from)
    if date_to:
        c_qs = c_qs.filter(payment_date__lte=date_to)
    labels = dict(FlipkartOtherCharge.CATEGORY_CHOICES)
    charges = list(c_qs.order_by("-payment_date", "-id"))
    cat = {}
    month_cat = {}
    for c in charges:
        a = cat.setdefault(c.category, {"category": c.category, "label": labels.get(c.category, c.category), "count": 0, "out": 0.0, "in": 0.0})
        amt = float(c.amount)
        a["count"] += 1
        a["out" if amt < 0 else "in"] += amt
        m = month_cat.setdefault(_month(c.payment_date), {"month": _month(c.payment_date)})
        m[c.category] = round(m.get(c.category, 0) + amt, 2)
    total_charges = sum(float(c.amount) for c in charges)
    ads_spend = -sum(float(c.amount) for c in charges if c.category == "ads")
    fines_total = -sum(float(c.amount) for c in charges if c.category == "fines")
    fines = {
        "count": len(charges),
        "total": round(total_charges, 2),
        "money_out": round(sum(float(c.amount) for c in charges if c.amount < 0), 2),
        "money_in": round(sum(float(c.amount) for c in charges if c.amount > 0), 2),
        "ads_spend": round(ads_spend, 2),
        "ads_pct_of_sales": round(ads_spend * 100.0 / total_sale, 1) if total_sale > 0 and ads_spend > 0 else None,
        "fines_total": round(fines_total, 2),
        "fines_pct_of_sales": round(fines_total * 100.0 / total_sale, 1) if total_sale > 0 and fines_total > 0 else None,
        "by_category": sorted(
            [{**a, "out": round(a["out"], 2), "in": round(a["in"], 2), "net": round(a["out"] + a["in"], 2)} for a in cat.values()],
            key=lambda x: x["net"]),
        "by_month": sorted(month_cat.values(), key=lambda x: x["month"]),
        "by_source": sorted(
            [{"label": k or "—", "total": round(t, 2), "count": n} for k, (t, n) in {
                s: (sum(float(c.amount) for c in charges if c.source_sheet == s), sum(1 for c in charges if c.source_sheet == s))
                for s in {c.source_sheet for c in charges}}.items()],
            key=lambda x: x["total"]),
        "biggest": [
            {"date": str(c.payment_date) if c.payment_date else None, "category": labels.get(c.category, c.category),
             "description": c.description, "reference": c.reference_id, "amount": float(c.amount)}
            for c in sorted(charges, key=lambda c: float(c.amount))[:10] if c.amount < 0
        ],
        "lines": [
            {"date": str(c.payment_date) if c.payment_date else None, "category": c.category,
             "category_label": labels.get(c.category, c.category), "sheet": c.source_sheet,
             "description": c.description, "reference": c.reference_id, "neft_id": c.neft_id, "amount": float(c.amount)}
            for c in charges[:200]
        ],
    }

    return Response({"fees": fees, "sla": sla, "fines": fines})


@api_view(["GET"])
def flipkart_sku_report(request, business_id):
    """
    Everything about one Flipkart SKU, for the "full report" window: its P&L for
    the period, the settlement components behind it, the month-by-month history
    (all time) and what the Orders report says (outcomes, SLA, reasons).
    """
    business = get_authorized_business(request, business_id)
    sku = (request.GET.get("sku") or "").strip()
    if not sku:
        return Response({"error": "sku is required"}, status=status.HTTP_400_BAD_REQUEST)
    date_from, date_to = _date_params(request)

    period = _compute_flipkart_profit_summary(business, date_from, date_to, sku=sku)
    alltime = _compute_flipkart_profit_summary(business, None, None, sku=sku)
    key = _sku_key(sku)
    row = next((r for r in period["sku_wise"] if _sku_key(r["sku_id"]) == key), None)
    parent = _flipkart_parent_link_map(business).get(sku)

    o_qs = FlipkartOrder.objects.filter(business=business, sku__iexact=sku)
    if date_from:
        o_qs = o_qs.filter(order_date__gte=date_from)
    if date_to:
        o_qs = o_qs.filter(order_date__lte=date_to)
    orders = list(o_qs)
    by_status = Counter(o.order_item_status or "UNKNOWN" for o in orders)
    live = [o for o in orders if o.order_item_status != STATUS_CANCELLED]

    return Response({
        "sku": sku,
        "parent_id": parent,
        "missing_price": key in {_sku_key(s) for s in period["missing_price_skus"]},
        "row": row,
        "totals": period["totals"],
        "status_breakdown": period["status_breakdown"],
        "payment_breakdown": period["payment_breakdown"],
        "fulfilment_breakdown": period["fulfilment_breakdown"],
        "monthly": alltime["monthly_trend"],
        "orders": {
            "total": len(orders),
            "units": sum(o.quantity for o in orders),
            "by_status": [{"status": s, "count": n} for s, n in by_status.most_common()],
            "delivered": by_status.get(STATUS_DELIVERED, 0),
            "returned": sum(by_status.get(s, 0) for s in _RETURN_STATUSES),
            "cancelled": by_status.get(STATUS_CANCELLED, 0),
            "dispatch_sla_breached": sum(1 for o in orders if o.dispatch_sla_breached),
            "delivery_sla_breached": sum(1 for o in orders if o.delivery_sla_breached),
            "live_items": len(live),
            "return_reasons": [{"reason": r, "count": n} for r, n in Counter(
                (o.return_reason + (f" — {o.return_sub_reason}" if o.return_sub_reason else "")) for o in orders if o.return_reason).most_common(8)],
            "cancel_reasons": [{"reason": r, "count": n} for r, n in Counter(
                o.cancellation_reason for o in orders if o.cancellation_reason).most_common(6)],
        },
    })


@api_view(["GET"])
def flipkart_sku_analysis(request, business_id):
    """
    One row per Flipkart SKU joining the three sources: profit (settlement
    payments costed against the SKU pricing table), the settlement components
    (sale, fees, taxes...), and the outcome counts from the Orders report.
    """
    business = get_authorized_business(request, business_id)
    date_from, date_to = _date_params(request)

    summary = _compute_flipkart_profit_summary(business, date_from, date_to)
    parent_link = _flipkart_parent_link_map(business)
    missing = {_sku_key(s) for s in summary["missing_price_skus"]}

    rows = {}

    def row_for(sku):
        return rows.setdefault(_sku_key(sku), {
            "sku_id": sku, "parent_id": parent_link.get(sku), "payment_items": 0,
            "delivered_count": 0, "return_count": 0, "gross_payout": 0.0, "total_cost": 0.0,
            "net_profit": 0.0, "sale_amount": 0.0, "settlement_value": 0.0,
            "marketplace_fee": 0.0, "taxes": 0.0, "refund": 0.0,
            "ordered_items": 0, "report_delivered": 0, "report_returned": 0, "report_cancelled": 0,
            "delivered_profit": 0.0, "return_loss": 0.0, "delivered_quantity": 0, "delivered_cost": 0.0,
            "one_unit_price": None,
            "is_loss_making": False, "required_adjustment_per_unit": None, "loss_all_from_returns": False,
        })

    for p in summary["sku_payment_list"]:
        r = row_for(p["sku_id"])
        r["payment_items"] += p["order_count"]
        for f in ("sale_amount", "settlement_value", "marketplace_fee", "taxes", "refund"):
            r[f] = round(r[f] + p[f], 2)
    for s in summary["sku_wise"]:
        r = row_for(s["sku_id"])
        r.update({k: s[k] for k in ("delivered_count", "return_count", "gross_payout", "total_cost", "net_profit",
                                    "is_loss_making", "required_adjustment_per_unit", "loss_all_from_returns",
                                    "delivered_profit", "return_loss", "delivered_quantity", "delivered_cost",
                                    "one_unit_price")})
        r["sku_id"] = s["sku_id"]

    orders_qs = FlipkartOrder.objects.filter(business=business)
    if date_from:
        orders_qs = orders_qs.filter(order_date__gte=date_from)
    if date_to:
        orders_qs = orders_qs.filter(order_date__lte=date_to)
    for sku, st, n in orders_qs.values_list("sku", "order_item_status").annotate(n=Count("id")):
        r = row_for(sku or "(blank SKU)")
        r["ordered_items"] += n
        if st == STATUS_DELIVERED:
            r["report_delivered"] += n
        elif st in _RETURN_STATUSES:
            r["report_returned"] += n
        elif st == STATUS_CANCELLED:
            r["report_cancelled"] += n

    skus = list(rows.values())
    for r in skus:
        sent = r["ordered_items"] - r["report_cancelled"]
        r["return_rate_pct"] = round(r["report_returned"] * 100.0 / sent, 1) if sent > 0 else None
        r["margin_pct"] = round(r["net_profit"] * 100.0 / r["sale_amount"], 1) if r["sale_amount"] > 0 else None
        r["missing_price"] = _sku_key(r["sku_id"]) in missing
    skus.sort(key=lambda r: r["net_profit"], reverse=True)

    return Response({
        "has_data": bool(skus),
        "totals": summary["totals"],
        "status_breakdown": summary["status_breakdown"],
        "missing_price_skus": summary["missing_price_skus"],
        "missing_price_count": summary["missing_price_count"],
        "skus": skus,
    })
