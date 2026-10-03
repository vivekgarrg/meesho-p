"""Parent-level Meesho stock sheet (SKU Analysis -> "take this off sale").

The sheet is pushed straight into Meesho's bulk inventory uploader, so the
format is not negotiable: these tests check the generated workbook against
Meesho's own export (the reference Inventory-Update-File in the repo root) and
round-trip it back through meesho_inventory_upload.
"""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.models import Business, Membership, User

from .models import FinalPrice, MeeshoInventory, ParentItemPrice

REFERENCE_SHEET = next(
    iter(sorted(Path(settings.BASE_DIR).parent.glob("Inventory-Update-File_*.xlsx"))), None,
)


def read_sheet(content):
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.worksheets[0]
    rows = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(1, ws.max_row + 1)
    ]
    return ws, rows


class StockSheetTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="owner", password="pw", role=User.ROLE_SUPER_ADMIN,
        )
        self.business = Business.objects.create(name="Test Business")
        Membership.objects.create(user=self.user, business=self.business)
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f"/api/business/{self.business.id}"

        self.parent = ParentItemPrice.objects.create(
            business=self.business, item_id="DIYA-THALI", item_price=Decimal("40"),
        )
        # Three SKUs on file under the parent. SKU-B is listed on Meesho under
        # two variations; SKU-C was never in an inventory upload.
        for sku in ("SKU-A", "SKU-B", "SKU-C"):
            FinalPrice.objects.create(business=self.business, sku_id=sku, parent=self.parent)
        # A SKU under no parent at all.
        FinalPrice.objects.create(business=self.business, sku_id="LONE-SKU")

        self._inv(1, "SKU-A", 588584002, 1110752420, 167, "Free Size", 1000)
        self._inv(2, "SKU-B", 588584001, 1110752419, 167, "Free Size", 20)
        self._inv(3, "SKU-B", 588584001, 1110752418, 168, "0-6 Months", 5)
        self._inv(4, "LONE-SKU", 588583998, 1110752416, 167, "Free Size", 7)

    def _inv(self, serial, style, catalog, product, variation_id, variation, system):
        return MeeshoInventory.objects.create(
            business=self.business, serial_no=serial,
            catalog_name=f"Catalog {catalog}", catalog_id=catalog,
            product_name=f"Product {product}", product_id=product,
            style_id=style, variation_id=variation_id, variation=variation,
            stock_type="ALL", system_stock_count=system,
        )

    def post_sheet(self, **body):
        return self.client.post(f"{self.base}/meesho-inventory/stock-sheet/", body, format="json")


class PreviewTests(StockSheetTestCase):
    def test_preview_reports_ready_and_missing(self):
        r = self.client.get(f"{self.base}/meesho-inventory/stock-sheet/preview/?parent_id=DIYA-THALI")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["parent"], "DIYA-THALI")
        self.assertEqual(r.data["sku_count"], 3)
        # SKU-A (1 variation) + SKU-B (2 variations) = 3 lines
        self.assertEqual(r.data["rows_ready"], 3)
        self.assertEqual(r.data["missing"], ["SKU-C"])

    def test_preview_exposes_catalog_ids(self):
        r = self.client.get(f"{self.base}/meesho-inventory/stock-sheet/preview/?parent_id=DIYA-THALI")
        by_sku = {m["sku_id"]: m for m in r.data["matched"]}
        self.assertEqual(by_sku["SKU-A"]["catalog_ids"], [588584002])
        self.assertEqual(len(by_sku["SKU-B"]["variations"]), 2)
        self.assertEqual(
            sorted(v["variation"] for v in by_sku["SKU-B"]["variations"]),
            ["0-6 Months", "Free Size"],
        )

    def test_a_sku_with_no_parent_is_its_own_group(self):
        """SKU Analysis groups by parent_item_id || sku_id, so this must work."""
        r = self.client.get(f"{self.base}/meesho-inventory/stock-sheet/preview/?parent_id=LONE-SKU")
        self.assertEqual(r.data["sku_count"], 1)
        self.assertEqual(r.data["rows_ready"], 1)

    def test_explicit_sku_ids_override_the_parent(self):
        r = self.client.get(f"{self.base}/meesho-inventory/stock-sheet/preview/?sku_ids=SKU-A,SKU-B")
        self.assertEqual(r.data["sku_count"], 2)
        self.assertEqual(r.data["rows_ready"], 3)


class SheetFormatTests(StockSheetTestCase):
    def test_zero_writes_a_real_zero_not_a_blank(self):
        """A blank YOUR STOCK COUNT means "no change" to Meesho, so a 0 that
        collapsed to blank would silently fail to stop the product."""
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        self.assertEqual(r.status_code, 200)
        _ws, rows = read_sheet(r.content)
        data = rows[2:]
        self.assertEqual(len(data), 3)
        for row in data:
            self.assertEqual(row[10], 0)
            self.assertNotEqual(row[10], "")

    def test_a_quantity_is_written_as_given(self):
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=25)
        _ws, rows = read_sheet(r.content)
        self.assertEqual([row[10] for row in rows[2:]], [25, 25, 25])

    def test_every_variation_becomes_its_own_line(self):
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        _ws, rows = read_sheet(r.content)
        styles = [row[5] for row in rows[2:]]
        self.assertEqual(sorted(styles), ["SKU-A", "SKU-B", "SKU-B"])
        self.assertEqual(r["X-Rows-Written"], "3")
        self.assertEqual(r["X-Skus-Written"], "2")
        self.assertEqual(r["X-Skus-Missing"], "1")

    def test_catalog_and_system_stock_are_carried_through(self):
        """The identifiers nobody can retype by hand, and the current count the
        seller needs in order to sanity-check what they are about to push."""
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        _ws, rows = read_sheet(r.content)
        row_a = next(row for row in rows[2:] if row[5] == "SKU-A")
        self.assertEqual(row_a[2], 588584002)   # CATALOG ID
        self.assertEqual(row_a[4], 1110752420)  # PRODUCT ID
        self.assertEqual(row_a[8], "ALL")       # STOCK type preserved
        self.assertEqual(row_a[9], 1000)        # SYSTEM STOCK COUNT

    def test_headers_match_meeshos_own_export_exactly(self):
        if REFERENCE_SHEET is None:
            self.skipTest("reference Inventory-Update-File not present")
        ref = openpyxl.load_workbook(REFERENCE_SHEET, read_only=True).worksheets[0]
        ref_h = [ref.cell(1, c).value for c in range(1, 12)]
        ref_d = [ref.cell(2, c).value for c in range(1, 12)]

        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        ws, rows = read_sheet(r.content)
        self.assertEqual(rows[0], ref_h)
        self.assertEqual(rows[1], ref_d)
        self.assertIn("fill this", ws.title.lower())

    def test_the_generated_sheet_re_uploads_cleanly(self):
        """The real proof of format: feed it back through the uploader that
        reads Meesho's own file."""
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        upload = SimpleUploadedFile(
            "stock.xlsx", r.content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        up = self.client.post(
            f"{self.base}/meesho-inventory/upload/", {"file": upload}, format="multipart",
        )
        self.assertIn(up.status_code, (200, 201), getattr(up, "data", up.content))
        self.assertEqual(up.data.get("skipped", 0), 0, up.data)
        # Every line matched an existing inventory row — nothing invented.
        self.assertEqual(up.data.get("created", 0), 0, up.data)
        self.assertEqual(up.data.get("updated", 0), 3, up.data)
        # And the 0 landed on the rows it was meant for.
        self.assertEqual(
            sorted(MeeshoInventory.objects.filter(business=self.business, seller_stock_count=0)
                   .values_list("style_id", flat=True)),
            ["SKU-A", "SKU-B", "SKU-B"],
        )


class ValidationTests(StockSheetTestCase):
    def test_a_missing_stock_count_is_rejected(self):
        r = self.post_sheet(parent_id="DIYA-THALI")
        self.assertEqual(r.status_code, 400)
        self.assertIn("stock quantity", r.data["error"].lower())

    def test_negative_and_absurd_counts_are_rejected(self):
        self.assertEqual(self.post_sheet(parent_id="DIYA-THALI", stock_count=-1).status_code, 400)
        self.assertEqual(self.post_sheet(parent_id="DIYA-THALI", stock_count=999999).status_code, 400)

    def test_a_non_numeric_count_is_rejected(self):
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count="lots")
        self.assertEqual(r.status_code, 400)

    def test_a_parent_with_nothing_in_inventory_explains_itself(self):
        bare = ParentItemPrice.objects.create(business=self.business, item_id="BARE")
        FinalPrice.objects.create(business=self.business, sku_id="NOT-ON-MEESHO", parent=bare)
        r = self.post_sheet(parent_id="BARE", stock_count=0)
        self.assertEqual(r.status_code, 400)
        self.assertIn("upload the inventory sheet", r.data["error"])
        self.assertEqual(r.data["missing"], ["NOT-ON-MEESHO"])

    def test_sku_casing_does_not_matter(self):
        """sku_id compares case-insensitively in the DB; style_id matching has
        to agree or a SKU silently drops out of the sheet."""
        MeeshoInventory.objects.filter(style_id="SKU-A").update(style_id="sku-a")
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["X-Rows-Written"], "3")


class IsolationTests(StockSheetTestCase):
    def test_another_business_cannot_build_a_sheet_for_this_parent(self):
        other = Business.objects.create(name="Someone Else")
        outsider = User.objects.create_user(username="x", password="pw", role=User.ROLE_BUSINESS_USER)
        Membership.objects.create(user=outsider, business=other)
        client = APIClient()
        client.force_authenticate(outsider)
        r = client.post(
            f"{self.base}/meesho-inventory/stock-sheet/",
            {"parent_id": "DIYA-THALI", "stock_count": 0}, format="json",
        )
        self.assertEqual(r.status_code, 404)

    def test_another_businesss_inventory_never_leaks_in(self):
        other = Business.objects.create(name="Other")
        MeeshoInventory.objects.create(
            business=other, serial_no=1, catalog_name="Theirs", catalog_id=999,
            product_name="Theirs", product_id=999, style_id="SKU-C",
            variation_id=1, variation="Free Size", stock_type="ALL", system_stock_count=50,
        )
        r = self.post_sheet(parent_id="DIYA-THALI", stock_count=0)
        _ws, rows = read_sheet(r.content)
        self.assertNotIn(999, [row[2] for row in rows[2:]])
        self.assertEqual(r["X-Skus-Missing"], "1")


class RealMeeshoFileTests(TestCase):
    """End to end on Meesho's actual export: upload it, group three of its real
    style ids under one parent, and pull that parent off sale."""

    def setUp(self):
        if REFERENCE_SHEET is None:
            self.skipTest("reference Inventory-Update-File not present")
        self.user = User.objects.create_user(
            username="owner", password="pw", role=User.ROLE_SUPER_ADMIN,
        )
        self.business = Business.objects.create(name="Real")
        Membership.objects.create(user=self.user, business=self.business)
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f"/api/business/{self.business.id}"

        with open(REFERENCE_SHEET, "rb") as fh:
            up = self.client.post(
                f"{self.base}/meesho-inventory/upload/",
                {"file": SimpleUploadedFile(REFERENCE_SHEET.name, fh.read())},
                format="multipart",
            )
        self.assertIn(up.status_code, (200, 201), getattr(up, "data", None))
        self.inventory_count = MeeshoInventory.objects.filter(business=self.business).count()
        self.assertGreater(self.inventory_count, 1000)

    def test_pull_a_real_parent_off_sale(self):
        styles = list(
            MeeshoInventory.objects.filter(business=self.business)
            .exclude(style_id="").values_list("style_id", flat=True).distinct()[:3]
        )
        parent = ParentItemPrice.objects.create(business=self.business, item_id="DIWALI-DIYA")
        for sku in styles:
            FinalPrice.objects.create(business=self.business, sku_id=sku, parent=parent)

        preview = self.client.get(
            f"{self.base}/meesho-inventory/stock-sheet/preview/?parent_id=DIWALI-DIYA"
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["sku_count"], 3)
        self.assertEqual(preview.data["missing"], [])
        self.assertGreaterEqual(preview.data["rows_ready"], 3)
        # Real catalog ids came through, not placeholders.
        for m in preview.data["matched"]:
            self.assertTrue(all(isinstance(c, int) and c > 0 for c in m["catalog_ids"]), m)

        sheet = self.client.post(
            f"{self.base}/meesho-inventory/stock-sheet/",
            {"parent_id": "DIWALI-DIYA", "stock_count": 0}, format="json",
        )
        self.assertEqual(sheet.status_code, 200)
        ws, rows = read_sheet(sheet.content)
        self.assertEqual(len(rows) - 2, int(sheet["X-Rows-Written"]))
        self.assertTrue(all(row[10] == 0 for row in rows[2:]))
        self.assertEqual(sorted({row[5] for row in rows[2:]}), sorted(styles))
        # Only this parent's SKUs — not all 1828 rows of the upload.
        self.assertLess(len(rows) - 2, self.inventory_count)

    def test_the_sheet_round_trips_into_the_same_rows(self):
        styles = list(
            MeeshoInventory.objects.filter(business=self.business)
            .exclude(style_id="").values_list("style_id", flat=True).distinct()[:2]
        )
        parent = ParentItemPrice.objects.create(business=self.business, item_id="P")
        for sku in styles:
            FinalPrice.objects.create(business=self.business, sku_id=sku, parent=parent)

        sheet = self.client.post(
            f"{self.base}/meesho-inventory/stock-sheet/",
            {"parent_id": "P", "stock_count": 0}, format="json",
        )
        back = self.client.post(
            f"{self.base}/meesho-inventory/upload/",
            {"file": SimpleUploadedFile("stop.xlsx", sheet.content)}, format="multipart",
        )
        self.assertIn(back.status_code, (200, 201), getattr(back, "data", b""))
        self.assertEqual(back.data.get("skipped", 0), 0, back.data)
        self.assertEqual(back.data.get("created", 0), 0, back.data)
        # Nothing outside this parent was touched, and the count is unchanged.
        self.assertEqual(
            MeeshoInventory.objects.filter(business=self.business).count(), self.inventory_count,
        )
        zeroed = set(
            MeeshoInventory.objects.filter(business=self.business, seller_stock_count=0)
            .values_list("style_id", flat=True)
        )
        self.assertEqual(zeroed, set(styles))


class TickedSubsetTests(StockSheetTestCase):
    """The Stock Update panel on a parent card sends parent_id AND the ticked
    SKUs, so that combination has its own rules."""

    def test_a_ticked_subset_is_honoured(self):
        r = self.post_sheet(parent_id="DIYA-THALI", sku_ids=["SKU-A"], stock_count=0)
        self.assertEqual(r.status_code, 200)
        _ws, rows = read_sheet(r.content)
        self.assertEqual([row[5] for row in rows[2:]], ["SKU-A"])
        self.assertEqual(r["X-Rows-Written"], "1")

    def test_a_sku_outside_the_parent_is_refused_not_zeroed(self):
        """A stale or mistyped id must never take some other product off sale."""
        other_parent = ParentItemPrice.objects.create(business=self.business, item_id="OTHER")
        FinalPrice.objects.create(business=self.business, sku_id="OTHER-SKU", parent=other_parent)
        self._inv(9, "OTHER-SKU", 777777, 888888, 167, "Free Size", 400)

        r = self.post_sheet(
            parent_id="DIYA-THALI", sku_ids=["SKU-A", "OTHER-SKU"], stock_count=0,
        )
        self.assertEqual(r.status_code, 200)
        _ws, rows = read_sheet(r.content)
        styles = [row[5] for row in rows[2:]]
        self.assertEqual(styles, ["SKU-A"])
        self.assertNotIn("OTHER-SKU", styles)
        self.assertNotIn(777777, [row[2] for row in rows[2:]])

    def test_every_variation_of_a_ticked_sku_is_still_written(self):
        r = self.post_sheet(parent_id="DIYA-THALI", sku_ids=["SKU-B"], stock_count=12)
        _ws, rows = read_sheet(r.content)
        self.assertEqual(len(rows) - 2, 2)
        self.assertTrue(all(row[10] == 12 for row in rows[2:]))

    def test_ticking_only_an_unmatched_sku_explains_itself(self):
        r = self.post_sheet(parent_id="DIYA-THALI", sku_ids=["SKU-C"], stock_count=0)
        self.assertEqual(r.status_code, 400)
        self.assertIn("upload the inventory sheet", r.data["error"])

    def test_sku_ids_without_a_parent_still_work(self):
        """SKU Analysis can group SKUs that have no ParentItemPrice at all."""
        r = self.post_sheet(sku_ids=["LONE-SKU"], stock_count=0)
        self.assertEqual(r.status_code, 200)
        _ws, rows = read_sheet(r.content)
        self.assertEqual([row[5] for row in rows[2:]], ["LONE-SKU"])
