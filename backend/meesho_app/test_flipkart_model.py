"""Flipkart as its own marketplace: Orders report, Settlement Report, and the P&L over both.

The workbooks built here copy the real exports' layout — the Settlement Report's
"Orders" sheet has a grouping row above the header and a blank row below it, the
Orders report prefixes item ids with "OI:" and wraps SKUs in quotes — so the
parsing is checked against what Flipkart really sends.
"""

import io
from decimal import Decimal

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.models import Business, Membership, User

from .models import FinalPrice, FlipkartOrder, FlipkartOrderPayment, FlipkartOtherCharge

SETTLEMENT_HEADERS = [
    "NEFT ID", "Neft Type", "Payment Date", "Bank Settlement Value (Rs.)\n= SUM(J:R)",
    "Order ID", "Order item ID", "Sale Amount (Rs.)", "Marketplace Fee (Rs.)", "Refund (Rs.)",
    "Order Date", "Fulfilment Type", "Seller SKU", "Quantity", "Return Type", "Item Return Status",
]


# These tests work their profit figures out by hand around other rules, so they
# run without the business-wide ₹6 per-order buffer (see test_default_buffer.py).
def setUpModule():
    from .models import BusinessCostSetting
    field = BusinessCostSetting._meta.get_field("default_buffer_per_order")
    field._saved_default, field.default = field.default, 0


def tearDownModule():
    from .models import BusinessCostSetting
    field = BusinessCostSetting._meta.get_field("default_buffer_per_order")
    field.default = field._saved_default


def settlement_row(neft, item, sale, settlement, sku="SKU-A", order="OD1", return_type="NA", refund=0.0):
    return {
        "NEFT ID": neft, "Neft Type": "Prepaid", "Payment Date": "2026-09-28",
        "Bank Settlement Value (Rs.)": settlement, "Order ID": order, "Order item ID": item,
        "Sale Amount (Rs.)": sale, "Marketplace Fee (Rs.)": 0.0, "Refund (Rs.)": refund,
        "Order Date": "2026-09-10", "Fulfilment Type": "seller_easy_ship", "Seller SKU": sku,
        "Quantity": 1.0, "Return Type": return_type, "Item Return Status": "NA",
    }


def settlement_workbook(order_rows, ads=(), fines=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Payment Details"])
    ws.append(SETTLEMENT_HEADERS)
    ws.append([None] * len(SETTLEMENT_HEADERS))
    for r in order_rows:
        ws.append([r.get(h.split("\n")[0].strip()) for h in SETTLEMENT_HEADERS])

    sheet = wb.create_sheet("Ads")
    sheet.append(["Payment Details"])
    sheet.append(["NEFT ID", "Payment Date", "Settlement Value (Rs.) = SUM(G:K)", None, "Type",
                  "Campaign / Transaction ID"])
    for neft, amount, kind, ref in ads:
        sheet.append([neft, "2026-09-21", amount, None, kind, ref])

    sheet = wb.create_sheet("Fines")
    sheet.append(["Payment Details"])
    sheet.append(["NEFT ID", "Payment Date", "Settlement Value(Rs.) ", None, "Transaction Name",
                  "Transaction Details", "Transaction ID"])
    for neft, amount, ref in fines:
        sheet.append([neft, "2026-09-21", amount, None, "Fines", "fines", ref])

    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile("settlement.xlsx", out.getvalue())


ORDER_HEADERS = ["order_item_id", "order_id", "order_date", "order_item_status", "sku", "product_title",
                 "quantity", "return_reason", "cancellation_reason", "dispatch_sla_breached",
                 "order_cancellation_date", "dispatched_date"]


def orders_workbook(rows):
    wb = openpyxl.Workbook()
    wb.active.title = "Help"
    ws = wb.create_sheet("Orders")
    ws.append(ORDER_HEADERS)
    for r in rows:
        ws.append([r.get(h, "") for h in ORDER_HEADERS])
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile("orders.xlsx", out.getvalue())


def order_row(item, status, sku="SKU-A", order="OD1", **extra):
    return {"order_item_id": f"OI:{item}", "order_id": order, "order_date": "2026-09-10 12:00:00",
            "order_item_status": status, "sku": f'"""SKU:{sku}"""', "product_title": '"""A Brass Kalash"""',
            "quantity": "1", **extra}


class FlipkartTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="owner", password="pw", role=User.ROLE_SUPER_ADMIN)
        self.business = Business.objects.create(name="B")
        Membership.objects.create(user=self.user, business=self.business)
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f"/api/business/{self.business.id}"
        # SKU-A costs 100 + 5% tax + 10 packaging = 115 to deliver one unit.
        FinalPrice.objects.create(business=self.business, sku_id="SKU-A", item_price=Decimal("100"),
                                  tax_percent=5, packaging_cost=Decimal("10"))

    def upload_settlement(self, **kw):
        return self.client.post(f"{self.base}/flipkart-profit/upload/",
                                {"file": settlement_workbook(**kw)}, format="multipart")

    def upload_orders(self, rows):
        return self.client.post(f"{self.base}/flipkart-orders/upload/",
                                {"file": orders_workbook(rows)}, format="multipart")


class OrdersReportUploadTests(FlipkartTestCase):
    def test_item_id_and_sku_are_cleaned_so_they_join_to_payments_and_pricing(self):
        r = self.upload_orders([order_row("438512", "DELIVERED")])
        self.assertEqual(r.status_code, 200)
        o = FlipkartOrder.objects.get()
        self.assertEqual(o.order_item_id, "438512")
        self.assertEqual(o.sku, "SKU-A")
        self.assertEqual(o.product_title, "A Brass Kalash")

    def test_blank_date_cells_are_stored_as_no_date(self):
        # Flipkart leaves cancellation/return/dispatch dates as empty strings.
        self.upload_orders([order_row("1", "DELIVERED")])
        o = FlipkartOrder.objects.get()
        self.assertIsNone(o.cancellation_date)
        self.assertIsNone(o.dispatched_date)

    def test_reupload_updates_status_in_place(self):
        self.upload_orders([order_row("1", "READY_TO_SHIP")])
        r = self.upload_orders([order_row("1", "DELIVERED")])
        self.assertEqual((r.data["created"], r.data["updated"]), (0, 1))
        self.assertEqual(FlipkartOrder.objects.get().order_item_status, "DELIVERED")

    def test_a_settlement_report_is_rejected_with_a_pointer_to_the_right_upload(self):
        r = self.client.post(f"{self.base}/flipkart-orders/upload/",
                             {"file": settlement_workbook([settlement_row("N1", "1", 200, 150)])},
                             format="multipart")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Settlement Report", r.data["error"])


class SettlementUploadTests(FlipkartTestCase):
    def test_ads_topups_and_fines_are_saved_but_zero_wallet_redeems_are_not(self):
        r = self.upload_settlement(
            order_rows=[settlement_row("N1", "1", 200, 150)],
            ads=[("N2", -3000.0, "topup", "SPWT1"), ("N3", 0.0, "redeem", "PCID1")],
            fines=[("N4", -240.0, "fine-1")],
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["other_charges_created"], 2)
        by_cat = {c.category: c.amount for c in FlipkartOtherCharge.objects.all()}
        self.assertEqual(by_cat, {"ads": Decimal("-3000.00"), "fines": Decimal("-240.00")})

    def test_reupload_does_not_duplicate_charges(self):
        kw = dict(order_rows=[settlement_row("N1", "1", 200, 150)], fines=[("N4", -240.0, "fine-1")])
        self.upload_settlement(**kw)
        r = self.upload_settlement(**kw)
        self.assertEqual((r.data["other_charges_created"], r.data["other_charges_updated"]), (0, 1))
        self.assertEqual(FlipkartOtherCharge.objects.count(), 1)


class ProfitTests(FlipkartTestCase):
    def test_delivered_profit_is_settlement_minus_sku_cost_then_charges_come_off(self):
        self.upload_settlement(
            order_rows=[settlement_row("N1", "1", 300, 200)],
            ads=[("N2", -50.0, "topup", "SPWT1")], fines=[("N3", -20.0, "f1")],
        )
        r = self.client.get(f"{self.base}/flipkart-overview/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["summary"]["totals"]["total_cost"], 115.0)
        self.assertEqual(r.data["summary"]["totals"]["net_profit"], 85.0)      # 200 - 115
        self.assertEqual(r.data["other_charges"]["total"], -70.0)
        self.assertEqual(r.data["net_profit"], 15.0)                           # 85 - 50 - 20

    def test_a_return_with_a_negative_settlement_is_a_loss_with_no_cost(self):
        self.upload_settlement(order_rows=[
            settlement_row("N1", "1", 300, -137, return_type="Customer Return", refund=-300)])
        totals = self.client.get(f"{self.base}/flipkart-overview/").data["summary"]["totals"]
        self.assertEqual((totals["total_cost"], totals["net_profit"]), (0.0, -137.0))

    def test_one_item_settled_as_sale_then_reversed_is_one_loss_not_two_rows(self):
        # Paid 200 on one NEFT, then the return took back 337 on the next: the
        # item is one unit that ended 137 down, and carries no purchase cost.
        self.upload_settlement(order_rows=[
            settlement_row("N1", "1", 300, 200),
            settlement_row("N2", "1", 0, -337, return_type="Customer Return"),
        ])
        totals = self.client.get(f"{self.base}/flipkart-overview/").data["summary"]["totals"]
        self.assertEqual((totals["order_count"], totals["total_cost"], totals["net_profit"]), (1, 0.0, -137.0))

    def test_an_unpriced_sku_is_left_out_of_profit(self):
        self.upload_settlement(order_rows=[settlement_row("N1", "1", 300, 200, sku="NO-PRICE")])
        data = self.client.get(f"{self.base}/flipkart-overview/").data
        self.assertEqual(data["summary"]["totals"]["net_profit"], 0.0)
        self.assertEqual(data["summary"]["missing_price_skus"], ["NO-PRICE"])


class GroupedPaymentsTests(FlipkartTestCase):
    def setUp(self):
        super().setUp()
        self.upload_orders([order_row("1", "DELIVERED", order="OD1"), order_row("2", "RETURNED", order="OD2")])
        self.upload_settlement(order_rows=[
            settlement_row("N1", "1", 300, 200, order="OD1"),
            settlement_row("N1", "2", 300, -137, order="OD2", return_type="Logistics Return"),
        ])

    def groups(self, query=""):
        return self.client.get(f"{self.base}/flipkart-payments/grouped/{query}").data

    def test_payments_are_grouped_by_order_id_with_profit_per_order(self):
        by_id = {g["order_id"]: g for g in self.groups()["groups"]}
        self.assertEqual(by_id["OD1"]["status"], "DELIVERED")
        self.assertEqual((by_id["OD1"]["settlement"], by_id["OD1"]["cost"], by_id["OD1"]["net_profit"]),
                         (200.0, 115.0, 85.0))
        self.assertEqual(by_id["OD2"]["status"], "RTO")
        self.assertEqual(by_id["OD2"]["net_profit"], -137.0)
        self.assertEqual(by_id["OD1"]["items"][0]["product_title"], "A Brass Kalash")

    def test_kpis_add_up_to_the_orders_shown(self):
        kpi = self.groups()["kpi"]
        self.assertEqual(kpi["total_groups"], 2)
        self.assertEqual(kpi["total_profit"], round(85.0 - 137.0, 2))

    def test_status_filter_narrows_the_list_but_not_the_kpis(self):
        data = self.groups("?status=RTO")
        self.assertEqual([g["order_id"] for g in data["groups"]], ["OD2"])
        self.assertEqual(data["kpi"]["total_groups"], 2)

    def test_search_matches_the_product_title_from_the_orders_report(self):
        self.assertEqual(self.groups("?search=Brass")["total"], 2)
        self.assertEqual(self.groups("?search=nothing-like-this")["total"], 0)


class OrdersListTests(FlipkartTestCase):
    def setUp(self):
        super().setUp()
        self.upload_orders([
            order_row("1", "DELIVERED", order="OD1"),                       # settled
            order_row("2", "DELIVERED", order="OD2"),                       # delivered, not paid yet
            order_row("3", "RETURNED", order="OD3", return_reason="QUALITY_ISSUE"),   # paid as a sale, now returned
            order_row("4", "CANCELLED", order="OD4", cancellation_reason="order_cancelled"),
        ])
        self.upload_settlement(order_rows=[
            settlement_row("N1", "1", 300, 200, order="OD1"),
            settlement_row("N1", "3", 300, 200, order="OD3"),
        ])

    def test_each_item_says_whether_it_has_been_settled(self):
        results = {r["order_item_id"]: r for r in
                   self.client.get(f"{self.base}/flipkart-orders/").data["results"]}
        self.assertTrue(results["1"]["settled"])
        self.assertEqual(results["1"]["settlement_value"], 200.0)
        self.assertFalse(results["2"]["settled"])
        self.assertIsNone(results["2"]["settlement_value"])

    def test_settled_filter(self):
        ids = lambda q: sorted(r["order_item_id"] for r in
                               self.client.get(f"{self.base}/flipkart-orders/?{q}").data["results"])
        self.assertEqual(ids("settled=yes"), ["1", "3"])
        self.assertEqual(ids("settled=no"), ["2", "4"])

    def test_analytics_flag_what_needs_attention(self):
        a = self.client.get(f"{self.base}/flipkart-orders/").data["analytics"]
        self.assertEqual((a["total_items"], a["delivered"], a["returned"], a["cancelled"]), (4, 2, 1, 1))
        self.assertEqual(a["delivered_unsettled"], 1)
        self.assertEqual(a["returned_after_settlement"], 1)
        self.assertEqual(a["cancel_rate_pct"], 25.0)
        self.assertEqual(a["return_rate_pct"], 33.3)    # 1 of the 3 items that weren't cancelled


class SkuAnalysisTests(FlipkartTestCase):
    def test_profit_settlement_and_order_outcomes_land_on_one_sku_row(self):
        self.upload_orders([order_row("1", "DELIVERED"), order_row("2", "RETURNED", order="OD2"),
                            order_row("3", "CANCELLED", order="OD3")])
        self.upload_settlement(order_rows=[settlement_row("N1", "1", 300, 200)])
        skus = self.client.get(f"{self.base}/flipkart-sku-analysis/").data["skus"]
        self.assertEqual(len(skus), 1)
        row = skus[0]
        self.assertEqual((row["net_profit"], row["sale_amount"]), (85.0, 300.0))
        self.assertEqual((row["ordered_items"], row["report_delivered"], row["report_returned"],
                          row["report_cancelled"]), (3, 1, 1, 1))
        self.assertEqual(row["return_rate_pct"], 50.0)


class BusinessIsolationTests(FlipkartTestCase):
    def test_another_business_sees_none_of_this_data(self):
        self.upload_orders([order_row("1", "DELIVERED")])
        self.upload_settlement(order_rows=[settlement_row("N1", "1", 300, 200)],
                               fines=[("N2", -20.0, "f1")])
        other = Business.objects.create(name="Other")
        Membership.objects.create(user=self.user, business=other)
        data = self.client.get(f"/api/business/{other.id}/flipkart-overview/").data
        self.assertFalse(data["has_data"])
        self.assertEqual(data["orders"]["total_items"], 0)
        self.assertEqual(data["other_charges"]["total"], 0.0)
        self.assertEqual(FlipkartOrderPayment.objects.filter(business=other).count(), 0)


class AvailableMonthsTests(FlipkartTestCase):
    def test_a_flipkart_only_business_still_gets_a_period_list(self):
        self.upload_orders([order_row("1", "DELIVERED")])                       # ordered 2026-09-10
        self.upload_settlement(order_rows=[settlement_row("N1", "1", 300, 200)])
        self.assertEqual(self.client.get(f"{self.base}/profit/available-months/").data, ["2026-09"])
